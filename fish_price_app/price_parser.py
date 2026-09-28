"""Read Trade81 price workbooks and normalize product names."""
from __future__ import annotations

from datetime import date, datetime
import re
import unicodedata
from typing import Any

from openpyxl import load_workbook


ALIASES = {
    "縞鯵": "シマアジ",
    "油甘魚": "ハマチ",
    "油金魚": "ハマチ",
    "油金": "ハマチ",
    "蜜柑鯛": "みかん鯛",
    "蜜甘鯛": "みかん鯛",
    "蜜柑鲷": "みかん鯛",
    # Trade81 order lists use both Japanese and Chinese names for farmed yellowtail.
    "養殖油甘魚": "ハマチ",
    "养殖油甘鱼": "ハマチ",
    "養殖はまち": "ハマチ",
    "養殖ハマチ": "ハマチ",
}


def is_ignored_product(value: Any) -> bool:
    """Return true for product lines assigned to other owners."""
    if value is None:
        return False
    text = unicodedata.normalize("NFKC", str(value)).strip()
    folded = text.casefold()
    return text.startswith("A") or any(token in folded for token in (
        "魚箱", "鱼箱", "fishbox", "fish box", "fish-box",
    ))


def _used_cells(ws) -> dict[tuple[int, int], Any]:
    # Some supplied workbooks have formatting extending to XFD. Only visit
    # cells that actually contain content.
    return {(c.row, c.column): c.value for c in ws._cells.values() if c.value is not None}


def _date_text(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    if value is None:
        return ""
    match = re.search(r"(20\d{2})[./年-]\s*(\d{1,2})[./月-]\s*(\d{1,2})", str(value))
    if match:
        return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
    return str(value).strip()


def _number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    raw = str(value).replace(",", "").strip()
    try:
        return float(raw)
    except ValueError:
        return None


def canonical_product_name(value: Any) -> str:
    """Standardize known aliases while retaining size, grade, and origin."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = re.sub(r"（[^）]*）|\([^)]*\)", "", text)
    text = re.sub(r"[×x＊*]\s*\d+(?:\.\d+)?\s*(?:尾|匹|個|枚|本|箱|p|パック)?", "", text)
    text = re.sub(r"\d+(?:\.\d+)?\s*(?:尾|匹|個|枚|本|箱)\s*$", "", text)
    # Some order rows append the ordered quantity as a bare number after a size,
    # e.g. "みかん鯛 1.8 kg 1". Keep the size and remove only that final count.
    text = re.sub(r"((?:\d+(?:\.\d+)?)\s*(?:kg|g|玉))\s+\d+\s*$", r"\1", text, flags=re.IGNORECASE)
    for old, new in sorted(ALIASES.items(), key=lambda item: len(item[0]), reverse=True):
        text = text.replace(old, new)
    # A description may list two synonymous names side by side, such as
    # "養殖ハマチ 養殖油甘魚". Collapse the duplicate after alias replacement.
    text = re.sub(r"(ハマチ)(?:\s*\1)+", r"\1", text)
    text = re.sub(r"\s+", " ", text)
    return text.strip(" ・,，;；")


def product_key(value: Any) -> str:
    """Conservative matching key; retain size/specification and remove spaces."""
    return re.sub(r"[\s・,，;；]+", "", canonical_product_name(value)).lower()


def _sheet_with_header(wb, row: int, expected: set[str]):
    for ws in wb.worksheets:
        cells = _used_cells(ws)
        row_values = {str(v).strip() for (r, _c), v in cells.items() if r == row}
        if expected.issubset(row_values):
            return ws, cells
    raise ValueError(f"找不到包含这些表头的工作表：{'、'.join(sorted(expected))}")


def read_price_template(file) -> dict[str, Any]:
    wb = load_workbook(file, data_only=True, read_only=False)
    ws, cells = _sheet_with_header(wb, 7, {"餐廳名稱", "產品名稱", "數量", "單位", "成本+5%", "总数"})
    report_date = _date_text(cells.get((5, 2)))
    factor = _number(cells.get((5, 10))) or 1.05
    rows = []
    for r in range(8, max((row for row, _col in cells), default=7) + 1):
        product = cells.get((r, 5))
        if not isinstance(product, str) or not product.strip():
            continue
        product = product.strip()
        if product.startswith("■"):
            continue
        # Sea urchin and fish-box lines are handled by other owners.
        if is_ignored_product(product):
            continue
        quote = _number(cells.get((r, 9)))
        rows.append({
            "product": product,
            "canonical_name": canonical_product_name(product),
            "key": product_key(product),
            "fishmonger_quote_jpy": quote,
            "source_row": r,
        })
    return {"date": report_date, "sheet": ws.title, "factor": factor, "rows": rows}


def read_order_template(file) -> dict[str, Any]:
    """Read a daily order workbook to count product order lines without prices."""
    wb = load_workbook(file, data_only=True, read_only=False)
    rows: list[dict[str, Any]] = []
    report_date = ""

    for ws in wb.worksheets:
        cells = _used_cells(ws)
        if not report_date:
            # Order workbooks usually place a title such as 新石下單 2026.9.28
            # near the top of the first sheet.
            for (r, _c), value in sorted(cells.items()):
                if r > 12:
                    continue
                parsed = _date_text(value)
                if parsed and re.fullmatch(r"20\d{2}-\d{2}-\d{2}", parsed):
                    report_date = parsed
                    break

        header_row = None
        product_col = None
        for r in range(1, min(max((row for row, _col in cells), default=1), 20) + 1):
            row_values = {str(v).strip() for (rr, _c), v in cells.items() if rr == r}
            if {"產品名稱", "數量", "單位"}.issubset(row_values):
                header_row = r
                product_col = next(c for (rr, c), v in cells.items() if rr == r and str(v).strip() == "產品名稱")
                break
        if header_row is None or product_col is None:
            continue

        for (r, c), value in cells.items():
            if r <= header_row or c != product_col or not isinstance(value, str):
                continue
            product = value.strip()
            if not product or product.startswith("■") or product in {"產品名稱", "产品名称"}:
                continue
            if is_ignored_product(product):
                continue
            canonical = canonical_product_name(product)
            key = product_key(product)
            if key:
                rows.append({"product": product, "canonical_name": canonical, "key": key})

    if not rows:
        raise ValueError("找不到每日下单模板中的商品行；需要有“產品名稱、數量、單位”表头。")
    if not report_date:
        raise ValueError("无法从每日下单模板标题读取日期。")
    return {"date": report_date, "rows": rows}


def group_price_rows(sheet: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], dict[str, list[float]]]:
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in sheet["rows"]:
        if row["key"] and row["fishmonger_quote_jpy"] is not None:
            grouped.setdefault(row["key"], []).append(row)
    unique: dict[str, dict[str, Any]] = {}
    conflicts: dict[str, list[float]] = {}
    for key, rows in grouped.items():
        prices = sorted({r["fishmonger_quote_jpy"] for r in rows})
        if len(prices) > 1:
            conflicts[key] = prices
        else:
            unique[key] = {**rows[0], "fishmonger_quote_jpy": prices[0]}
    return unique, conflicts
