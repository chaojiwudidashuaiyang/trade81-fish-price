"""SQLite for local development and PostgreSQL for the shared cloud version."""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
import sqlite3
from typing import Any

from openpyxl import Workbook, load_workbook


_DATABASE_URL: str | None = None


def configure_database(database_url: str | None) -> None:
    global _DATABASE_URL
    _DATABASE_URL = database_url or None


class _Connection:
    def __init__(self, raw, postgres: bool):
        self.raw = raw
        self.postgres = postgres

    def execute(self, query: str, params=()):
        if self.postgres:
            query = query.replace("?", "%s")
        return self.raw.execute(query, params)

    def commit(self):
        self.raw.commit()

    def rollback(self):
        self.raw.rollback()

    def close(self):
        self.raw.close()


def _local_path() -> Path:
    root = os.environ.get("LOCALAPPDATA")
    folder = Path(root) / "Trade81FishPrice" if root else Path.home() / ".trade81_fish_price"
    folder.mkdir(parents=True, exist_ok=True)
    return folder / "prices.db"


def connect() -> _Connection:
    postgres = bool(_DATABASE_URL)
    if postgres:
        try:
            import psycopg
            from psycopg.rows import dict_row
        except ImportError as exc:
            raise RuntimeError("云端数据库驱动未安装；请确认 requirements.txt 已包含 psycopg。") from exc
        raw = psycopg.connect(_DATABASE_URL, row_factory=dict_row, connect_timeout=15, prepare_threshold=None)
        db = _Connection(raw, True)
        db.execute("""
            CREATE TABLE IF NOT EXISTS price_catalog (
                product_key TEXT PRIMARY KEY,
                product_name TEXT NOT NULL,
                supplier_quote_jpy DOUBLE PRECISION NOT NULL,
                source_date TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                source TEXT NOT NULL,
                updated_by TEXT NOT NULL DEFAULT ''
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS price_history (
                id BIGSERIAL PRIMARY KEY,
                product_key TEXT NOT NULL,
                product_name TEXT NOT NULL,
                supplier_quote_jpy DOUBLE PRECISION NOT NULL,
                source_date TEXT NOT NULL,
                updated_at TEXT NOT NULL,
                source TEXT NOT NULL,
                updated_by TEXT NOT NULL DEFAULT ''
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS order_daily_counts (
                order_date TEXT NOT NULL,
                product_key TEXT NOT NULL,
                product_name TEXT NOT NULL,
                order_count BIGINT NOT NULL,
                PRIMARY KEY(order_date, product_key)
            )
        """)
        db.execute("""
            CREATE TABLE IF NOT EXISTS product_aliases (
                alias_key TEXT PRIMARY KEY,
                alias_name TEXT NOT NULL,
                canonical_key TEXT NOT NULL,
                canonical_name TEXT NOT NULL
            )
        """)
        db.execute("ALTER TABLE price_catalog ADD COLUMN IF NOT EXISTS updated_by TEXT NOT NULL DEFAULT ''")
        db.execute("ALTER TABLE price_history ADD COLUMN IF NOT EXISTS updated_by TEXT NOT NULL DEFAULT ''")
        db.commit()
        return db

    raw = sqlite3.connect(_local_path())
    raw.row_factory = sqlite3.Row
    db = _Connection(raw, False)
    db.execute("""
        CREATE TABLE IF NOT EXISTS price_catalog (
            product_key TEXT PRIMARY KEY,
            product_name TEXT NOT NULL,
            supplier_quote_jpy REAL NOT NULL,
            source_date TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            source TEXT NOT NULL,
            updated_by TEXT NOT NULL DEFAULT ''
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS price_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product_key TEXT NOT NULL,
            product_name TEXT NOT NULL,
            supplier_quote_jpy REAL NOT NULL,
            source_date TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            source TEXT NOT NULL,
            updated_by TEXT NOT NULL DEFAULT ''
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS order_daily_counts (
            order_date TEXT NOT NULL,
            product_key TEXT NOT NULL,
            product_name TEXT NOT NULL,
            order_count INTEGER NOT NULL,
            PRIMARY KEY(order_date, product_key)
        )
    """)
    db.execute("""
        CREATE TABLE IF NOT EXISTS product_aliases (
            alias_key TEXT PRIMARY KEY,
            alias_name TEXT NOT NULL,
            canonical_key TEXT NOT NULL,
            canonical_name TEXT NOT NULL
        )
    """)
    # Add the audit field to a local database created by an earlier app version.
    for table in ("price_catalog", "price_history"):
        columns = {row[1] for row in db.execute(f"PRAGMA table_info({table})").fetchall()}
        if "updated_by" not in columns:
            db.execute(f"ALTER TABLE {table} ADD COLUMN updated_by TEXT NOT NULL DEFAULT ''")
    db.commit()
    return db


@contextmanager
def db_session():
    db = connect()
    try:
        yield db
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def list_prices() -> list[dict[str, Any]]:
    with db_session() as db:
        price_rows = [dict(row) for row in db.execute("SELECT * FROM price_catalog").fetchall()]
        count_rows = [dict(row) for row in db.execute("SELECT product_key, SUM(order_count) AS order_count FROM order_daily_counts GROUP BY product_key").fetchall()]
        alias_rows = [dict(row) for row in db.execute("SELECT alias_key, alias_name, canonical_key, canonical_name FROM product_aliases").fetchall()]

    alias_map = {row["alias_key"]: row for row in alias_rows}
    groups: dict[str, dict[str, Any]] = {}

    def target_for(key: str, fallback_name: str = "") -> tuple[str, str]:
        alias = alias_map.get(key)
        return (alias["canonical_key"], alias["canonical_name"]) if alias else (key, fallback_name)

    for row in price_rows:
        canonical_key, canonical_name = target_for(row["product_key"], row["product_name"])
        bucket = groups.setdefault(canonical_key, {
            "product_key": canonical_key, "product_name": canonical_name,
            "member_keys": [], "order_count": 0, "_prices": [], "aliases": set(),
        })
        bucket["member_keys"].append(row["product_key"])
        bucket["_prices"].append(row)
    for row in count_rows:
        canonical_key, _canonical_name = target_for(row["product_key"])
        if canonical_key in groups:
            groups[canonical_key]["order_count"] += int(row["order_count"] or 0)
    for row in alias_rows:
        if row["canonical_key"] in groups and row["alias_key"] != row["canonical_key"]:
            groups[row["canonical_key"]]["aliases"].add(row["alias_name"])

    result = []
    for bucket in groups.values():
        newest = max(bucket.pop("_prices"), key=lambda row: (row["source_date"], row["updated_at"], row["product_key"] == bucket["product_key"]))
        bucket.update({key: newest[key] for key in ("supplier_quote_jpy", "source_date", "updated_at", "source", "updated_by")})
        bucket["member_keys"] = sorted(set(bucket["member_keys"]))
        bucket["aliases"] = sorted(bucket["aliases"])
        result.append(bucket)
    return sorted(result, key=lambda row: (-row["order_count"], row["product_name"].casefold()))


def list_product_groups() -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("""
            SELECT alias_key, alias_name, canonical_key, canonical_name
            FROM product_aliases ORDER BY canonical_name, alias_name
        """).fetchall()
    groups: dict[str, dict[str, Any]] = {}
    for raw in rows:
        row = dict(raw)
        group = groups.setdefault(row["canonical_key"], {
            "canonical_key": row["canonical_key"],
            "canonical_name": row["canonical_name"],
            "aliases": [],
        })
        if row["alias_key"] != row["canonical_key"] and row["alias_name"] != row["canonical_name"]:
            group["aliases"].append(row["alias_name"])
    return list(groups.values())


def resolve_sheet_products(sheet: dict[str, Any]) -> dict[str, Any]:
    """Attach custom display keys to rows, retaining raw keys for price/order storage."""
    with db_session() as db:
        aliases = {
            row["alias_key"]: dict(row)
            for row in db.execute("SELECT alias_key, alias_name, canonical_key, canonical_name FROM product_aliases").fetchall()
        }
    resolved = {**sheet, "rows": []}
    for original in sheet.get("rows", []):
        row = dict(original)
        row.setdefault("raw_key", row.get("key", ""))
        row.setdefault("raw_name", row.get("canonical_name") or row.get("product") or row.get("key", ""))
        alias = aliases.get(row["raw_key"])
        if alias:
            row["key"] = alias["canonical_key"]
            row["canonical_name"] = alias["canonical_name"]
        resolved["rows"].append(row)
    return resolved


def save_product_group(primary_name: str, alias_names: list[str]) -> None:
    """Add or combine user-defined names under one display name without deleting source rows."""
    from price_parser import canonical_product_name, product_key

    canonical_name = canonical_product_name(primary_name)
    canonical_key = product_key(canonical_name)
    if not canonical_key:
        raise ValueError("请填写主要显示品名。")

    submitted: dict[str, str] = {canonical_key: canonical_name}
    for name in alias_names:
        display_name = canonical_product_name(name)
        key = product_key(display_name)
        if key:
            submitted[key] = display_name
    if len(submitted) < 2:
        raise ValueError("请至少填写一个需要合并的其他品名。")

    with db_session() as db:
        existing = [dict(row) for row in db.execute("SELECT alias_key, alias_name, canonical_key, canonical_name FROM product_aliases").fetchall()]
        submitted_keys = set(submitted)
        source_groups = {canonical_key}
        source_groups.update(row["canonical_key"] for row in existing if row["alias_key"] in submitted_keys)
        members = dict(submitted)
        for row in existing:
            if row["canonical_key"] in source_groups:
                members[row["alias_key"]] = row["alias_name"]

        # Existing groups touched by any selected name are reassigned as a whole.
        db.execute("DELETE FROM product_aliases")
        for alias_key, alias_name in members.items():
            db.execute("""
                INSERT INTO product_aliases(alias_key, alias_name, canonical_key, canonical_name)
                VALUES (?, ?, ?, ?)
            """, (alias_key, alias_name, canonical_key, canonical_name))
        # Preserve untouched groups.
        for row in existing:
            if row["canonical_key"] not in source_groups and row["alias_key"] not in members:
                db.execute("""
                    INSERT INTO product_aliases(alias_key, alias_name, canonical_key, canonical_name)
                    VALUES (?, ?, ?, ?)
                """, (row["alias_key"], row["alias_name"], row["canonical_key"], row["canonical_name"]))


def delete_product_group(canonical_key: str) -> None:
    """Remove a display grouping; original prices, history, and counts remain intact."""
    with db_session() as db:
        db.execute("DELETE FROM product_aliases WHERE canonical_key=?", (canonical_key,))


def record_order_counts(sheet: dict[str, Any]) -> int:
    """Replace one date's counts with occurrences of each item in the uploaded sheet.

    Each product row is one order occurrence (the quantity column is not used).
    Re-uploading a corrected workbook for the same date replaces that date's counts,
    so retries do not inflate cumulative totals.
    """
    order_date = str(sheet.get("date") or "")
    if not order_date:
        raise ValueError("订单统计需要有有效日期。")

    counts: dict[str, dict[str, Any]] = {}
    for item in sheet.get("rows", []):
        key = str(item.get("raw_key") or item.get("key") or "").strip()
        if not key:
            continue
        if key not in counts:
            counts[key] = {
                "name": item.get("raw_name") or item.get("canonical_name") or item.get("product") or key,
                "count": 0,
            }
        counts[key]["count"] += 1

    with db_session() as db:
        db.execute("DELETE FROM order_daily_counts WHERE order_date=?", (order_date,))
        for key, item in counts.items():
            db.execute("""
                INSERT INTO order_daily_counts(order_date, product_key, product_name, order_count)
                VALUES (?, ?, ?, ?)
            """, (order_date, key, item["name"], item["count"]))
    return sum(item["count"] for item in counts.values())


def list_history(limit: int | None = 30) -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("SELECT * FROM price_history ORDER BY id DESC").fetchall()
        aliases = {
            row["alias_key"]: dict(row)
            for row in db.execute("SELECT alias_key, canonical_key, canonical_name FROM product_aliases").fetchall()
        }
    result = []
    seen = set()
    for raw in rows:
        row = dict(raw)
        alias = aliases.get(row["product_key"])
        if alias:
            row["product_key"] = alias["canonical_key"]
            row["product_name"] = alias["canonical_name"]
        signature = (
            row["product_key"], row["source_date"], row["supplier_quote_jpy"],
            row["source"], row["updated_by"],
        )
        if signature in seen:
            continue
        seen.add(signature)
        result.append(row)
        if limit is not None and len(result) >= limit:
            break
    return result


def _record_change(db: _Connection, key: str, name: str, quote: float,
                   source_date: str, source: str, now: str, updated_by: str) -> None:
    db.execute("""
        INSERT INTO price_history(product_key, product_name, supplier_quote_jpy, source_date, updated_at, source, updated_by)
        VALUES (?, ?, ?, ?, ?, ?, ?)
    """, (key, name, quote, source_date, now, source, updated_by))


def apply_sheet(sheet: dict[str, Any], updated_by: str = "") -> dict[str, int]:
    """Import a dated sheet; newer source dates supersede older/manual values."""
    changed = skipped = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db_session() as db:
        for item in sheet["rows"]:
            key = item.get("raw_key") or item["key"]
            quote = item["fishmonger_quote_jpy"]
            if not key or quote is None:
                continue
            old = db.execute("SELECT * FROM price_catalog WHERE product_key=?", (key,)).fetchone()
            if old and sheet["date"] < old["source_date"]:
                skipped += 1
                continue
            if old and sheet["date"] == old["source_date"] and old["source"] == "手动修改":
                skipped += 1
                continue
            if old and sheet["date"] == old["source_date"] and old["source"] == "价格表" and old["supplier_quote_jpy"] == quote:
                skipped += 1
                continue
            name = item.get("raw_name") or item.get("canonical_name") or item["product"]
            db.execute("""
                INSERT INTO price_catalog(product_key, product_name, supplier_quote_jpy, source_date, updated_at, source, updated_by)
                VALUES (?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(product_key) DO UPDATE SET
                    product_name=excluded.product_name,
                    supplier_quote_jpy=excluded.supplier_quote_jpy,
                    source_date=excluded.source_date,
                    updated_at=excluded.updated_at,
                    source=excluded.source,
                    updated_by=excluded.updated_by
            """, (key, name, quote, sheet["date"], now, "价格表", updated_by))
            _record_change(db, key, name, quote, sheet["date"], "价格表", now, updated_by)
            changed += 1
    return {"changed": changed, "skipped": skipped}


def manual_update(key: str, quote: float, updated_by: str = "", member_keys: list[str] | None = None) -> None:
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with db_session() as db:
        keys = sorted(set(member_keys or [key]))
        for raw_key in keys:
            old = db.execute("SELECT * FROM price_catalog WHERE product_key=?", (raw_key,)).fetchone()
            if not old:
                continue
            name = old["product_name"]
            db.execute("""
                UPDATE price_catalog SET supplier_quote_jpy=?, updated_at=?, source='手动修改', updated_by=?
                WHERE product_key=?
            """, (quote, now, updated_by, raw_key))
            _record_change(db, raw_key, name, quote, old["source_date"], "手动修改", now, updated_by)


def export_backup_xlsx() -> bytes:
    wb = Workbook()
    catalog = wb.active
    catalog.title = "当前价格"
    catalog.append(["product_key", "product_name", "supplier_quote_jpy", "source_date", "updated_at", "source", "updated_by"])
    for row in list_raw_prices():
        catalog.append([row.get(k) for k in ("product_key", "product_name", "supplier_quote_jpy", "source_date", "updated_at", "source", "updated_by")])
    history = wb.create_sheet("更新历史")
    history.append(["product_key", "product_name", "supplier_quote_jpy", "source_date", "updated_at", "source", "updated_by"])
    for row in reversed(list_raw_history()):
        history.append([row.get(k) for k in ("product_key", "product_name", "supplier_quote_jpy", "source_date", "updated_at", "source", "updated_by")])
    counts = wb.create_sheet("每日下单统计")
    counts.append(["order_date", "product_key", "product_name", "order_count"])
    for row in list_order_counts():
        counts.append([row.get(k) for k in ("order_date", "product_key", "product_name", "order_count")])
    groups = wb.create_sheet("品名归并")
    groups.append(["alias_key", "alias_name", "canonical_key", "canonical_name"])
    for row in list_alias_rows():
        groups.append([row.get(k) for k in ("alias_key", "alias_name", "canonical_key", "canonical_name")])
    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()


def list_raw_prices() -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("SELECT * FROM price_catalog ORDER BY LOWER(product_name)").fetchall()
    return [dict(row) for row in rows]


def list_raw_history() -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("SELECT * FROM price_history ORDER BY id").fetchall()
    return [dict(row) for row in rows]


def list_order_counts() -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("""
            SELECT order_date, product_key, product_name, order_count
            FROM order_daily_counts
            ORDER BY order_date, product_key
        """).fetchall()
    return [dict(row) for row in rows]


def list_alias_rows() -> list[dict[str, Any]]:
    with db_session() as db:
        rows = db.execute("""
            SELECT alias_key, alias_name, canonical_key, canonical_name
            FROM product_aliases ORDER BY canonical_name, alias_name
        """).fetchall()
    return [dict(row) for row in rows]


def import_backup_xlsx(file) -> int:
    wb = load_workbook(file, data_only=True, read_only=True)
    if "当前价格" not in wb.sheetnames:
        raise ValueError("迁移文件里没有“当前价格”工作表。")
    catalog_ws = wb["当前价格"]
    headers = [cell.value for cell in catalog_ws[1]]
    records = [dict(zip(headers, row)) for row in catalog_ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in row)]
    if not records:
        raise ValueError("迁移文件中没有价格记录。")
    with db_session() as db:
        existing = db.execute("SELECT COUNT(*) AS n FROM price_catalog").fetchone()["n"]
        if existing:
            raise ValueError("云端价格库已有数据。为了避免覆盖，请仅在空价格库中导入迁移文件。")
        for row in records:
            db.execute("""
                INSERT INTO price_catalog(product_key, product_name, supplier_quote_jpy, source_date, updated_at, source, updated_by)
                VALUES (?, ?, ?, ?, ?, ?, ?)
            """, tuple(row.get(k) or "" for k in ("product_key", "product_name")) + (
                float(row.get("supplier_quote_jpy") or 0), str(row.get("source_date") or ""),
                str(row.get("updated_at") or ""), str(row.get("source") or "迁移导入"), str(row.get("updated_by") or ""),
            ))
        if "更新历史" in wb.sheetnames:
            history_ws = wb["更新历史"]
            history_headers = [cell.value for cell in history_ws[1]]
            for values in history_ws.iter_rows(min_row=2, values_only=True):
                if not any(v is not None for v in values):
                    continue
                row = dict(zip(history_headers, values))
                db.execute("""
                    INSERT INTO price_history(product_key, product_name, supplier_quote_jpy, source_date, updated_at, source, updated_by)
                    VALUES (?, ?, ?, ?, ?, ?, ?)
                """, (
                    str(row.get("product_key") or ""), str(row.get("product_name") or ""),
                    float(row.get("supplier_quote_jpy") or 0), str(row.get("source_date") or ""),
                    str(row.get("updated_at") or ""), str(row.get("source") or "迁移导入"), str(row.get("updated_by") or ""),
                ))
        if "每日下单统计" in wb.sheetnames:
            counts_ws = wb["每日下单统计"]
            count_headers = [cell.value for cell in counts_ws[1]]
            for values in counts_ws.iter_rows(min_row=2, values_only=True):
                if not any(v is not None for v in values):
                    continue
                row = dict(zip(count_headers, values))
                db.execute("""
                    INSERT INTO order_daily_counts(order_date, product_key, product_name, order_count)
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(order_date, product_key) DO UPDATE SET
                        product_name=excluded.product_name,
                        order_count=excluded.order_count
                """, (
                    str(row.get("order_date") or ""), str(row.get("product_key") or ""),
                    str(row.get("product_name") or ""), int(row.get("order_count") or 0),
                ))
        if "品名归并" in wb.sheetnames:
            groups_ws = wb["品名归并"]
            group_headers = [cell.value for cell in groups_ws[1]]
            for values in groups_ws.iter_rows(min_row=2, values_only=True):
                if not any(v is not None for v in values):
                    continue
                row = dict(zip(group_headers, values))
                db.execute("""
                    INSERT INTO product_aliases(alias_key, alias_name, canonical_key, canonical_name)
                    VALUES (?, ?, ?, ?)
                    ON CONFLICT(alias_key) DO UPDATE SET
                        alias_name=excluded.alias_name,
                        canonical_key=excluded.canonical_key,
                        canonical_name=excluded.canonical_name
                """, tuple(str(row.get(k) or "") for k in ("alias_key", "alias_name", "canonical_key", "canonical_name")))
    return len(records)
